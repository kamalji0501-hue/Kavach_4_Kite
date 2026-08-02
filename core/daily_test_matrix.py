"""Excel test matrix for daily UAT execution — append rows per run."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.batman_mode import daily_test_execution_dir, workspace_root
from core.uat_daily_tests import TEST_CATALOG, CaseResult

IST = ZoneInfo("Asia/Kolkata")

MATRIX_DIR = daily_test_execution_dir()
MATRIX_FILE = MATRIX_DIR / "test_matrix.xlsx"
SHEET_MATRIX = "Matrix"
LATEST_MD = MATRIX_DIR / "LATEST_RUN.md"

HEADERS = (
    "date_ist",
    "time_ist",
    "case_id",
    "category",
    "name",
    "status",
    "duration_ms",
    "detail",
    "metrics_json",
    "run_id",
)


def _catalog_lookup() -> dict[str, tuple[str, str]]:
    return {c.case_id: (c.category, c.name) for c in TEST_CATALOG}


def ensure_workbook(path: Path | None = None) -> Path:
    path = path or MATRIX_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file():
        return path
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_MATRIX
    ws.append(list(HEADERS))
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    wb.save(path)
    return path


def append_results(
    results: list[CaseResult],
    *,
    run_id: str | None = None,
    path: Path | None = None,
) -> Path:
    path = ensure_workbook(path)
    run_id = run_id or datetime.now(tz=IST).strftime("%Y%m%d_%H%M%S")
    now = datetime.now(tz=IST)
    date_s = now.strftime("%Y-%m-%d")
    time_s = now.strftime("%H:%M:%S")
    lookup = _catalog_lookup()

    wb = load_workbook(path)
    if SHEET_MATRIX not in wb.sheetnames:
        ws = wb.create_sheet(SHEET_MATRIX)
        ws.append(list(HEADERS))
    ws = wb[SHEET_MATRIX]

    daily_name = f"Daily_{date_s}"
    if daily_name not in wb.sheetnames:
        dws = wb.create_sheet(daily_name)
        dws.append(list(HEADERS))
        for col in range(1, len(HEADERS) + 1):
            dws.cell(row=1, column=col).font = Font(bold=True)
    dws = wb[daily_name]

    for r in results:
        cat, name = lookup.get(r.case_id, ("", ""))
        if r.skipped:
            status = "SKIP"
        elif r.ok:
            status = "PASS"
        else:
            status = "FAIL"
        metrics_json = ""
        if r.metrics:
            import json

            metrics_json = json.dumps(r.metrics, default=str)
        row = [
            date_s,
            time_s,
            r.case_id,
            cat,
            name,
            status,
            round(r.duration_ms, 1),
            (r.detail or "")[:500],
            metrics_json[:500],
            run_id,
        ]
        ws.append(row)
        dws.append(row)

    for sheet in (ws, dws):
        for col in range(1, len(HEADERS) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18

    wb.save(path)
    return path


def write_latest_summary(
    results: list[CaseResult],
    *,
    matrix_path: Path,
    run_id: str,
    elapsed_s: float,
    log_path: Path | None = None,
) -> Path:
    MATRIX_DIR.mkdir(parents=True, exist_ok=True)
    passed = sum(1 for r in results if r.ok and not r.skipped)
    failed = sum(1 for r in results if not r.ok and not r.skipped)
    skipped = sum(1 for r in results if r.skipped)
    lines = [
        "# Latest UAT daily test run",
        "",
        f"- **Run ID:** `{run_id}`",
        f"- **When (IST):** {datetime.now(tz=IST).strftime('%Y-%m-%d %H:%M:%S')}",
        f"- **Duration:** {elapsed_s:.1f}s",
        f"- **PASS:** {passed} · **FAIL:** {failed} · **SKIP:** {skipped}",
        "",
    ]
    if log_path:
        lines.append(f"- **Suite log:** `{log_path.resolve()}`")
        lines.append("")
    lines.extend(
        [
        f"## Test matrix (Excel)",
        "",
        f"**File:** `{matrix_path.resolve()}`",
        "",
        "| Case | Status | Detail |",
        "|------|--------|--------|",
        ]
    )
    for r in results:
        st = "SKIP" if r.skipped else ("PASS" if r.ok else "FAIL")
        lines.append(f"| {r.case_id} | {st} | {(r.detail or '')[:80]} |")
    lines.extend(
        [
            "",
            "## Agent command",
            "",
            "```powershell",
            ".venv\\Scripts\\python.exe scripts\\run_uat_daily_test_suite.py",
            "```",
            "",
        ]
    )
    LATEST_MD.write_text("\n".join(lines), encoding="utf-8")
    return LATEST_MD
