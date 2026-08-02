"""Export incident knowledge registry to a detailed multi-sheet Excel workbook."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.incident_knowledge import IncidentRecord, load_registry, registry_path

_IST = ZoneInfo("Asia/Kolkata")
_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = _ROOT / "docs" / "incidents" / "incident_registry.xlsx"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _join(items: list[str] | None) -> str:
    if not items:
        return ""
    return "; ".join(str(x) for x in items if x)


def _commentary_line(rec: IncidentRecord) -> str:
    """Human-readable narrative: issue → root cause → fix date → solution."""
    observed = rec.date_first_observed or "unknown"
    fixed = rec.date_resolved or "open"
    issue = rec.symptoms or rec.title
    parts = [
        f"Observed {observed}: {issue}",
        f"Root cause: {rec.root_cause or 'under investigation'}",
    ]
    if rec.fix_implemented:
        parts.append(f"Fixed {fixed}: {rec.fix_implemented}")
    if rec.validation_performed:
        parts.append(f"Validated: {_join(rec.validation_performed)}")
    if rec.lessons_learned:
        parts.append(f"Lesson: {_join(rec.lessons_learned)}")
    return " | ".join(parts)


def _fixture_expiry_context() -> list[dict[str, str]]:
    """UAT Sensibull book expiry + auto-roll commentary for Fixture_Expiry sheet."""
    fixture_path = _ROOT / "uat" / "deployed_positions" / "positions.json"
    rows: list[dict[str, str]] = [
        {"field": "fixture_path", "value": str(fixture_path.relative_to(_ROOT)).replace("\\", "/")},
    ]
    if not fixture_path.is_file():
        rows.append({"field": "status", "value": "MISSING — ingest screenshot via /register"})
        return rows

    try:
        raw = json.loads(fixture_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        rows.append({"field": "status", "value": f"READ ERROR: {exc}"})
        return rows

    label = str(raw.get("expiry_label") or "")
    expiry_s = str(raw.get("expiry_date") or "")
    captured = str(raw.get("captured_at") or "")
    legs = raw.get("legs") or []
    rows.extend(
        [
            {"field": "expiry_label_in_fixture", "value": label},
            {"field": "expiry_date_in_fixture", "value": expiry_s},
            {"field": "captured_at", "value": captured},
            {"field": "leg_count", "value": str(len(legs))},
        ]
    )

    resolved = ""
    rolled = "no"
    try:
        from datetime import datetime as dt

        from backtest_engine.resolver.instrument_master import resolve_fixture_trading_expiry

        fixture_expiry = dt.strptime(expiry_s, "%Y-%m-%d").date()
        resolved_date, did_roll = resolve_fixture_trading_expiry(fixture_expiry, legs)
        resolved = resolved_date.isoformat()
        rolled = "yes" if did_roll else "no"
        rows.extend(
            [
                {"field": "resolved_trading_expiry", "value": resolved},
                {"field": "auto_roll_active", "value": rolled},
            ]
        )
    except Exception as exc:
        rows.append({"field": "resolved_trading_expiry", "value": f"resolve error: {exc}"})

    fix_rec = next((r for r in load_registry() if r.incident_id == "INC-2026-016"), None)
    fix_date = fix_rec.date_resolved if fix_rec else "2026-06-25"
    rows.extend(
        [
            {"field": "code_fix_date", "value": fix_date},
            {"field": "linked_incident", "value": "INC-2026-016"},
            {
                "field": "what_was_the_issue",
                "value": (
                    fix_rec.symptoms
                    if fix_rec
                    else "Expired weekly expiry in positions.json blocked ShadowBroker"
                ),
            },
            {
                "field": "how_we_solved_it",
                "value": (
                    fix_rec.fix_implemented
                    if fix_rec
                    else "resolve_fixture_trading_expiry() rolls to nearest valid weekly"
                ),
            },
            {
                "field": "commentary",
                "value": (
                    f"Fixture file still shows {label or expiry_s} (captured {captured[:10]}). "
                    f"Code fix dated {fix_date} auto-rolls trading to {resolved or 'nearest weekly'} "
                    f"when file expiry is past. Operator should refresh Sensibull screenshot after "
                    f"each weekly expiry so label matches live book."
                ),
            },
            {
                "field": "operator_action",
                "value": "Refresh uat/deployed_positions screenshot after weekly NIFTY expiry",
            },
        ]
    )
    return rows


def _style_header_row(ws, ncol: int) -> None:
    for col in range(1, ncol + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP


def _autosize_columns(ws, max_width: int = 60) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, min(ws.max_row + 1, 200)):
            val = ws.cell(row=row, column=col_idx).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), max_width))
        ws.column_dimensions[letter].width = max(12, min(max_len + 2, max_width))


def _record_to_full_row(rec: IncidentRecord) -> dict[str, Any]:
    return {
        "incident_id": rec.incident_id,
        "title": rec.title,
        "category": rec.category,
        "subsystem": rec.subsystem,
        "robot_affected": rec.robot_affected,
        "date_first_observed": rec.date_first_observed,
        "date_resolved": rec.date_resolved,
        "lifecycle": rec.lifecycle,
        "severity": rec.severity,
        "priority": rec.priority,
        "environment": rec.environment,
        "frequency": rec.frequency,
        "symptoms": rec.symptoms,
        "business_impact": rec.business_impact,
        "technical_impact": rec.technical_impact,
        "trading_impact": rec.trading_impact,
        "root_cause": rec.root_cause,
        "contributing_factors": _join(rec.contributing_factors),
        "investigation": rec.investigation,
        "files_modules": _join(rec.files_modules),
        "logs_referenced": _join(rec.logs_referenced),
        "fix_implemented": rec.fix_implemented,
        "fix_rationale": rec.fix_rationale,
        "regression_risk": rec.regression_risk,
        "validation_performed": _join(rec.validation_performed),
        "stress_testing": _join(rec.stress_testing),
        "validation_cycles": rec.validation_cycles,
        "remaining_risks": _join(rec.remaining_risks),
        "related_incidents": _join(rec.related_incidents),
        "runtime_incident_ids": _join(rec.runtime_incident_ids),
        "preventive_measures": _join(rec.preventive_measures),
        "lessons_learned": _join(rec.lessons_learned),
        "future_recommendations": _join(rec.future_recommendations),
        "owner": rec.owner,
        "last_updated": rec.last_updated,
        "tags": _join(rec.tags),
    }


def _record_to_commentary_row(rec: IncidentRecord) -> dict[str, str]:
    return {
        "incident_id": rec.incident_id,
        "title": rec.title,
        "date_observed": rec.date_first_observed,
        "date_fixed": rec.date_resolved or "(open)",
        "status": rec.lifecycle,
        "severity": rec.severity,
        "priority": rec.priority,
        "robot": rec.robot_affected,
        "category": rec.category,
        "subsystem": rec.subsystem,
        "what_was_the_issue": rec.symptoms or rec.title,
        "root_cause": rec.root_cause,
        "how_we_solved_it": rec.fix_implemented,
        "validation_evidence": _join(rec.validation_performed),
        "related_incidents": _join(rec.related_incidents),
        "commentary": _commentary_line(rec),
    }


def export_incident_registry_xlsx(
    output_path: Path | None = None,
    *,
    registry: Path | None = None,
) -> Path:
    """Write detailed incident workbook (Summary, Commentary, All_Issues, Fixture_Expiry, …)."""
    out = output_path or DEFAULT_XLSX
    out.parent.mkdir(parents=True, exist_ok=True)
    items = sorted(load_registry(registry), key=lambda r: r.incident_id)
    now = datetime.now(tz=_IST)

    wb = Workbook()
    # --- Summary ---
    ws_sum = wb.active
    ws_sum.title = "Summary"
    dash_open = sum(1 for r in items if r.lifecycle != "closed")
    dash_closed = len(items) - dash_open
    reg_p = registry_path(registry)
    try:
        reg_display = str(reg_p.relative_to(_ROOT)).replace("\\", "/")
    except ValueError:
        reg_display = str(reg_p)
    summary_rows = [
        ("exported_at_ist", now.strftime("%Y-%m-%d %H:%M:%S")),
        ("registry_path", reg_display),
        ("total_incidents", str(len(items))),
        ("closed", str(dash_closed)),
        ("open", str(dash_open)),
        ("critical", str(sum(1 for r in items if r.severity == "critical"))),
        ("high", str(sum(1 for r in items if r.severity == "high"))),
        ("fixture_expiry_sheet", "See Fixture_Expiry tab for UAT book expiry + fix date commentary"),
        ("commentary_sheet", "See Commentary tab — issue, root cause, fix date, solution per incident"),
    ]
    ws_sum.append(["metric", "value"])
    for row in summary_rows:
        ws_sum.append(list(row))
    _style_header_row(ws_sum, 2)

    # --- Commentary (operator-friendly) ---
    ws_com = wb.create_sheet("Commentary")
    com_headers = list(_record_to_commentary_row(items[0]).keys()) if items else []
    if items:
        ws_com.append(com_headers)
        _style_header_row(ws_com, len(com_headers))
        for rec in items:
            row = _record_to_commentary_row(rec)
            ws_com.append([row[h] for h in com_headers])
            for col in range(1, len(com_headers) + 1):
                ws_com.cell(row=ws_com.max_row, column=col).alignment = _WRAP

    # --- All_Issues (full grid) ---
    ws_all = wb.create_sheet("All_Issues")
    if items:
        full_headers = list(_record_to_full_row(items[0]).keys())
        ws_all.append(full_headers)
        _style_header_row(ws_all, len(full_headers))
        for rec in items:
            row = _record_to_full_row(rec)
            ws_all.append([row[h] for h in full_headers])
            for col in range(1, len(full_headers) + 1):
                ws_all.cell(row=ws_all.max_row, column=col).alignment = _WRAP

    # --- Fixture_Expiry ---
    ws_fix = wb.create_sheet("Fixture_Expiry")
    ws_fix.append(["field", "value"])
    _style_header_row(ws_fix, 2)
    for row in _fixture_expiry_context():
        ws_fix.append([row["field"], row["value"]])
        ws_fix.cell(row=ws_fix.max_row, column=2).alignment = _WRAP

    # --- By_Robot ---
    ws_robot = wb.create_sheet("By_Robot")
    ws_robot.append(
        ["robot", "incident_id", "date_observed", "date_fixed", "severity", "title", "status", "commentary"]
    )
    _style_header_row(ws_robot, 8)
    for rec in items:
        robot = rec.robot_affected or "(system)"
        ws_robot.append(
            [
                robot,
                rec.incident_id,
                rec.date_first_observed,
                rec.date_resolved,
                rec.severity,
                rec.title,
                rec.lifecycle,
                _commentary_line(rec),
            ]
        )
        ws_robot.cell(row=ws_robot.max_row, column=8).alignment = _WRAP

    # --- Related_Incidents ---
    ws_rel = wb.create_sheet("Related_Incidents")
    ws_rel.append(["incident_id", "related_incident_id", "title", "date_fixed", "commentary"])
    _style_header_row(ws_rel, 5)
    id_map = {r.incident_id: r for r in items}
    for rec in items:
        for rel_id in rec.related_incidents:
            other = id_map.get(rel_id)
            ws_rel.append(
                [
                    rec.incident_id,
                    rel_id,
                    other.title if other else "",
                    other.date_resolved if other else "",
                    _commentary_line(other) if other else "",
                ]
            )

    # --- Lessons_Learned ---
    ws_less = wb.create_sheet("Lessons_Learned")
    ws_less.append(["incident_id", "date_fixed", "lesson"])
    _style_header_row(ws_less, 3)
    for rec in items:
        for lesson in rec.lessons_learned:
            ws_less.append([rec.incident_id, rec.date_resolved, lesson])

    # --- Open_Only ---
    ws_open = wb.create_sheet("Open_Only")
    open_headers = com_headers if items else []
    if open_headers:
        ws_open.append(open_headers)
        _style_header_row(ws_open, len(open_headers))
        for rec in items:
            if rec.lifecycle == "closed":
                continue
            row = _record_to_commentary_row(rec)
            ws_open.append([row[h] for h in open_headers])

    for ws in wb.worksheets:
        _autosize_columns(ws)

    wb.save(out)

    # Dated archive copy
    archive_dir = _ROOT / "data" / "analytics" / "incidents"
    archive_dir.mkdir(parents=True, exist_ok=True)
    archive = archive_dir / f"incident_registry_{now.strftime('%Y%m%d')}.xlsx"
    if archive.resolve() != out.resolve():
        wb.save(archive)

    return out


def default_xlsx_path() -> Path:
    return DEFAULT_XLSX
