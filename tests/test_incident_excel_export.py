"""Tests for incident registry Excel export."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from core.incident_excel_export import export_incident_registry_xlsx
from core.incident_seed import build_seed_incidents
from core.incident_knowledge import save_registry


def test_export_incident_xlsx_sheets(tmp_path: Path) -> None:
    reg = tmp_path / "registry.json"
    save_registry(build_seed_incidents(), reg)
    out = tmp_path / "incident_registry.xlsx"
    path = export_incident_registry_xlsx(out, registry=reg)
    assert path.is_file()
    wb = load_workbook(path, read_only=True)
    names = set(wb.sheetnames)
    assert "Summary" in names
    assert "Commentary" in names
    assert "All_Issues" in names
    assert "Fixture_Expiry" in names
    assert "By_Robot" in names
    wb.close()


def test_commentary_has_fix_dates(tmp_path: Path) -> None:
    reg = tmp_path / "registry.json"
    save_registry(build_seed_incidents(), reg)
    out = tmp_path / "incident_registry.xlsx"
    export_incident_registry_xlsx(out, registry=reg)
    wb = load_workbook(out, read_only=True)
    ws = wb["Commentary"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "date_observed" in headers
    assert "date_fixed" in headers
    assert "what_was_the_issue" in headers
    assert "how_we_solved_it" in headers
    assert "commentary" in headers
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) >= 20
    wb.close()


def test_fixture_expiry_sheet(tmp_path: Path, monkeypatch) -> None:
    reg = tmp_path / "registry.json"
    save_registry(build_seed_incidents(), reg)
    out = tmp_path / "incident_registry.xlsx"
    export_incident_registry_xlsx(out, registry=reg)
    wb = load_workbook(out, read_only=True)
    ws = wb["Fixture_Expiry"]
    fields = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert "expiry_date_in_fixture" in fields or "status" in fields
    assert "commentary" in fields or "fixture_path" in fields
    wb.close()
